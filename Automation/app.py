"""
app.py  –  EDW vs Power BI Data Reconciliation Tool
Run with:  streamlit run app.py

Output workbook sheets:
  Source | Target | Data_Rec | Summary | Column_Diff (if unmatched columns exist)
"""

import io
import time

import pandas as pd
import streamlit as st

from difflib import SequenceMatcher
from reconcile_engine import reconcile


def _similarity(a: str, b: str) -> float:
    """Return 0-1 similarity score between two strings (case-insensitive)."""
    return SequenceMatcher(None, a.lower(), b.lower()).ratio()


def _suggest_mapping(src_only: list, tgt_only: list, threshold: float = 0.6) -> dict:
    """
    Auto-suggest mappings between unmatched columns using string similarity.
    Returns {tgt_col: src_col} for pairs with similarity >= threshold.
    Each column can only be used once.
    """
    suggestions = {}
    used_src = set()
    used_tgt = set()
    # Build all pairs sorted by score descending
    pairs = []
    for sc in src_only:
        for tc in tgt_only:
            pairs.append((sc, tc, _similarity(sc, tc)))
    pairs.sort(key=lambda x: x[2], reverse=True)
    for sc, tc, score in pairs:
        if score >= threshold and sc not in used_src and tc not in used_tgt:
            suggestions[tc] = sc   # tgt → src
            used_src.add(sc)
            used_tgt.add(tc)
    return suggestions

# ── Page config ───────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Data Rec Tool – EDW vs Power BI",
    page_icon="🔍",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ── Custom CSS ────────────────────────────────────────────────────────────────
st.markdown(
    """
    <style>
    html, body, [class*="css"] { font-family: 'Inter', 'Segoe UI', sans-serif; }

    .banner {
        background: linear-gradient(120deg, #1a3a5c 0%, #0d6e8a 100%);
        color: white; padding: 1.6rem 2rem 1.2rem;
        border-radius: 0.75rem; margin-bottom: 1.5rem;
    }
    .banner h1 { margin: 0; font-size: 1.8rem; font-weight: 700; letter-spacing: -0.3px; }
    .banner p  { margin: 0.3rem 0 0; font-size: 0.95rem; opacity: 0.85; }

    .step-card {
        background: #f8fafd; border: 1px solid #d5e3f0;
        border-left: 4px solid #0d6e8a; border-radius: 0.5rem;
        padding: 1rem 1.2rem; margin-bottom: 1rem;
    }
    .step-card h4 { margin: 0 0 0.5rem; color: #1a3a5c; font-size: 1rem; }

    .metric-row { display: flex; flex-wrap: wrap; gap: 0.75rem; margin-top: 1rem; }
    .metric-pill {
        background: white; border: 1px solid #d5e3f0; border-radius: 0.5rem;
        padding: 0.5rem 1rem; min-width: 160px; text-align: center;
        box-shadow: 0 1px 4px rgba(0,0,0,.06);
    }
    .metric-pill .val { font-size: 1.4rem; font-weight: 700; color: #1a3a5c; }
    .metric-pill .lbl { font-size: 0.72rem; color: #6b7280; text-transform: uppercase; letter-spacing: .5px; }
    .metric-pill.warn .val { color: #9c6500; }
    .metric-pill.ok   .val { color: #276221; }
    .metric-pill.err  .val { color: #9c0006; }

    hr { margin: 1.5rem 0; border-color: #e5ecf3; }
    .timer-badge {
        display: inline-block; background: #e0f0ff; color: #1a3a5c;
        font-weight: 700; font-size: 0.9rem;
        padding: 0.25rem 0.75rem; border-radius: 1rem; margin-top: 0.5rem;
    }
    .col-diff-box {
        background: #fff8e1; border: 1px solid #f4b942;
        border-left: 4px solid #f4b942; border-radius: 0.5rem;
        padding: 0.75rem 1rem; margin-top: 0.5rem;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ── Banner ────────────────────────────────────────────────────────────────────
st.markdown(
    """
    <div class="banner">
      <h1>🔍 Data Reconciliation Tool</h1>
      <p>Automate EDW vs Power BI comparison · Generates Source · Target · Data_Rec · Summary · Column_Diff<br>
         <small>⚡ Vectorized processing · Date-format normalisation · Positional duplicate matching</small>
      </p>
    </div>
    """,
    unsafe_allow_html=True,
)

# ══════════════════════════════════════════════════════════════════════════════
# Step 1 – Upload
# ══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="step-card"><h4>① Upload Workbook</h4>', unsafe_allow_html=True)
uploaded = st.file_uploader(
    "Upload your Excel workbook (.xlsx / .xls)",
    type=["xlsx", "xls"],
    help="Must contain at least 2 sheets: Source (EDW) and Target (Power BI).",
)
st.markdown("</div>", unsafe_allow_html=True)

if not uploaded:
    st.info("⬆️  Upload a workbook above to get started.", icon="📂")
    st.stop()

try:
    xls = pd.ExcelFile(uploaded)
    sheet_names = xls.sheet_names
except Exception as e:
    st.error(f"❌ Could not read workbook: {e}")
    st.stop()

if len(sheet_names) < 2:
    st.error("❌ The workbook must contain at least 2 sheets.")
    st.stop()

st.success(f"✅ Workbook loaded — {len(sheet_names)} sheet(s): {', '.join(sheet_names)}", icon="📗")

# ══════════════════════════════════════════════════════════════════════════════
# Step 2 – Select sheets
# ══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="step-card"><h4>② Select Sheets</h4>', unsafe_allow_html=True)
c1, c2 = st.columns(2)
with c1:
    src_sheet = st.selectbox("Source sheet (EDW)", sheet_names, index=0)
with c2:
    rem = [s for s in sheet_names if s != src_sheet]
    tgt_sheet = st.selectbox("Target sheet (Power BI)", rem, index=0)
if src_sheet == tgt_sheet:
    st.error("❌ Source and Target sheets must be different.")
    st.stop()
st.markdown("</div>", unsafe_allow_html=True)

# ── Load data ─────────────────────────────────────────────────────────────────
@st.cache_data(show_spinner="Reading sheets…")
def load_sheets(file_bytes: bytes, src: str, tgt: str):
    buf = io.BytesIO(file_bytes)
    src_df = pd.read_excel(buf, sheet_name=src, dtype=str)
    buf.seek(0)
    tgt_df = pd.read_excel(buf, sheet_name=tgt, dtype=str)
    return src_df, tgt_df

file_bytes = uploaded.getvalue()
try:
    source_df, target_df = load_sheets(file_bytes, src_sheet, tgt_sheet)
except Exception as e:
    st.error(f"❌ Failed to load sheets: {e}")
    st.stop()

if source_df.empty or target_df.empty:
    st.error("❌ One or both selected sheets are empty.")
    st.stop()

st.info(
    f"📊 **Source**: {len(source_df):,} rows × {len(source_df.columns)} columns  |  "
    f"**Target**: {len(target_df):,} rows × {len(target_df.columns)} columns",
    icon="ℹ️",
)

# ── Column analysis ───────────────────────────────────────────────────────────
src_col_set = set(source_df.columns)
tgt_col_set = set(target_df.columns)
common_cols  = [c for c in source_df.columns if c in tgt_col_set]
src_only     = [c for c in source_df.columns if c not in tgt_col_set]
tgt_only     = [c for c in target_df.columns if c not in src_col_set]

with st.expander("👁 Preview loaded data", expanded=False):
    t_s, t_t = st.tabs([f"Source – {src_sheet}", f"Target – {tgt_sheet}"])
    with t_s:
        st.caption(f"{len(source_df):,} rows × {len(source_df.columns)} columns")
        st.dataframe(source_df.head(100), use_container_width=True)
    with t_t:
        st.caption(f"{len(target_df):,} rows × {len(target_df.columns)} columns")
        st.dataframe(target_df.head(100), use_container_width=True)

# ══════════════════════════════════════════════════════════════════════════════
# Step 2.5 – Column mapping (only shown when there are unmatched columns)
# ══════════════════════════════════════════════════════════════════════════════

# col_mapping accumulates user decisions: {tgt_col → src_col}
# We store it in session_state so it survives reruns.
if "col_mapping" not in st.session_state:
    st.session_state["col_mapping"] = {}

# Reset mapping whenever the loaded file / sheets change
_file_key = (uploaded.name, src_sheet, tgt_sheet)
if st.session_state.get("_file_key") != _file_key:
    st.session_state["_file_key"] = _file_key
    st.session_state["col_mapping"] = {}

# Only show this step if there are unmatched columns
if src_only or tgt_only:
    st.markdown(
        '<div class="step-card"><h4>② Map Unmatched Columns <span style="font-size:0.8rem;'
        'font-weight:400;color:#6b7280;">(optional)</span></h4>',
        unsafe_allow_html=True,
    )
    st.caption(
        "Some column names differ between Source and Target. "
        "Map them here so they compare together in Data_Rec. "
        "Columns left as **(Leave unmatched)** will appear in the **Column_Diff** sheet instead."
    )

    # Auto-suggest on first load for this file
    _auto_key = "_auto_mapped_" + str(hash(_file_key))
    if _auto_key not in st.session_state:
        suggestions = _suggest_mapping(src_only, tgt_only, threshold=0.55)
        st.session_state["col_mapping"].update(suggestions)
        st.session_state[_auto_key] = True

    # Build reverse: src_col → tgt_col (to block a src col from being mapped twice)
    _reverse = {v: k for k, v in st.session_state["col_mapping"].items()}

    for src_col in src_only:
        # Which tgt col is this src_col currently mapped from (if any)?
        currently_mapped_from = _reverse.get(src_col)   # tgt_col name or None

        # Available tgt options: unmatched tgt cols that are not already claimed by another src col
        claimed_by_others = {
            tc for tc, sc in st.session_state["col_mapping"].items() if sc != src_col
        }
        available_tgt = [t for t in tgt_only if t not in claimed_by_others]

        # Build display options
        options = ["(Leave unmatched)"] + available_tgt
        current_idx = 0
        if currently_mapped_from and currently_mapped_from in available_tgt:
            current_idx = available_tgt.index(currently_mapped_from) + 1

        col_l, col_r = st.columns([2, 3])
        with col_l:
            st.markdown(
                f"Map source column &nbsp;<code style='color:#0d6e8a;background:#e8f4f8;"
                f"padding:2px 6px;border-radius:4px;'>{src_col}</code>&nbsp; to target column:",
                unsafe_allow_html=True,
            )
        with col_r:
            choice = st.selectbox(
                label=f"tgt_for_{src_col}",
                options=options,
                index=current_idx,
                label_visibility="collapsed",
                key=f"map_{src_col}",
            )

        # Update mapping
        if choice == "(Leave unmatched)":
            # Remove any existing mapping for this src col
            st.session_state["col_mapping"] = {
                tc: sc for tc, sc in st.session_state["col_mapping"].items()
                if sc != src_col
            }
            st.caption(
                f"<span style='color:#6b7280;font-size:0.8rem;'>"
                f"`{src_col}` will remain unmatched and show as extra column in Source.</span>",
                unsafe_allow_html=True,
            )
        else:
            st.session_state["col_mapping"][choice] = src_col   # tgt_col → src_col
            st.caption(
                f"<span style='color:#276221;font-size:0.8rem;'>"
                f"✓ Mapped `{choice}` → `{src_col}`</span>",
                unsafe_allow_html=True,
            )

    # Show remaining unmatched tgt cols (not mapped to anything)
    mapped_tgt_cols = set(st.session_state["col_mapping"].keys())
    unmatched_tgt = [t for t in tgt_only if t not in mapped_tgt_cols]
    if unmatched_tgt:
        st.caption(
            "Target columns with no mapping (will go to Column_Diff): "
            + ", ".join(f"`{c}`" for c in unmatched_tgt)
        )

    st.markdown("</div>", unsafe_allow_html=True)

# After mapping, recompute what is truly common vs unmatched
_current_mapping = st.session_state.get("col_mapping", {})   # {tgt_col: src_col}
# Virtual target df (column names after rename) for computing common/unmatched
_tgt_virtual_cols = set()
for c in target_df.columns:
    _tgt_virtual_cols.add(_current_mapping.get(c, c))

src_col_set_after = set(source_df.columns)
common_cols  = [c for c in source_df.columns if c in _tgt_virtual_cols]
src_only_final = [c for c in source_df.columns if c not in _tgt_virtual_cols]
# tgt_only_final: tgt cols that are not mapped and not already in source
_mapped_tgt_names = set(_current_mapping.keys())
tgt_only_final = [
    c for c in target_df.columns
    if c not in _mapped_tgt_names          # not mapped away
    and c not in src_col_set_after         # not already common
]

# Summary banner: what will be in Data_Rec vs Column_Diff
if src_only_final or tgt_only_final:
    parts = []
    if src_only_final:
        parts.append(
            f"**{len(src_only_final)} column(s) only in Source** → Column_Diff: "
            + ", ".join(f"`{c}`" for c in src_only_final)
        )
    if tgt_only_final:
        parts.append(
            f"**{len(tgt_only_final)} column(s) only in Target** → Column_Diff: "
            + ", ".join(f"`{c}`" for c in tgt_only_final)
        )
    st.markdown(
        '<div class="col-diff-box">⚠️ ' + " &nbsp;|&nbsp; ".join(parts) + "<br>"
        "These will NOT appear in Data_Rec. They will be listed in the <b>Column_Diff</b> sheet.</div>",
        unsafe_allow_html=True,
    )
elif not src_only and not tgt_only:
    pass   # all columns matched exactly — no banner needed
else:
    st.success("✅ All unmatched columns have been mapped. Data_Rec will include all columns.", icon="✅")

# ══════════════════════════════════════════════════════════════════════════════
# Step 3 – Key columns
# ══════════════════════════════════════════════════════════════════════════════
st.markdown('<div class="step-card"><h4>③ Select Key Columns</h4>', unsafe_allow_html=True)
st.caption(
    "Select columns that together uniquely identify each record. "
    "Key columns must exist in **both** sheets."
)
key_cols = st.multiselect(
    "Key columns",
    options=common_cols,   # only common columns can be keys
    help="e.g. Policy Number, Claim Number + Effective Date.",
)

if key_cols:
    def _key_str(df, cols):
        """Build a normalized key string from `cols` with safe, case-insensitive lookup.

        If a requested column name is not present in `df` (e.g. different casing or trimmed),
        we try to find a case-insensitive match. If still not found, an empty string series
        is used so the function does not raise KeyError.
        """
        # helper to resolve a requested column name to an actual column in df
        def _resolve(col_name: str):
            if col_name in df.columns:
                return col_name
            target = col_name.strip().upper()
            for c in df.columns:
                if str(c).strip().upper() == target:
                    return c
            return None

        parts = []
        for c in cols:
            actual = _resolve(c)
            if actual is None:
                parts.append(pd.Series([""] * len(df), index=df.index))
            else:
                parts.append(df[actual].fillna("").astype(str).str.strip().str.upper())

        if len(parts) == 1:
            return parts[0]

        r = parts[0]
        for p in parts[1:]:
            r = r + "|||" + p
        return r

    src_dup = int(_key_str(source_df, key_cols).duplicated().sum())
    tgt_dup = int(_key_str(target_df, key_cols).duplicated().sum())
    if src_dup > 0:
        st.warning(
            f"⚠️ Source has **{src_dup:,} duplicate key combination(s)**. "
            "Rows will be matched positionally within each group — no cross-join duplication."
        )
    if tgt_dup > 0:
        st.warning(
            f"⚠️ Target has **{tgt_dup:,} duplicate key combination(s)**. "
            "Rows will be matched positionally within each group — no cross-join duplication."
        )

st.markdown("</div>", unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# Step 4 – Financial columns
# ══════════════════════════════════════════════════════════════════════════════
remaining_cols = [c for c in common_cols if c not in key_cols]

st.markdown('<div class="step-card"><h4>④ Select Financial Columns</h4>', unsafe_allow_html=True)
st.caption(
    "Financial columns → compared as **Target − Source** (numeric).  "
    "All other non-key common columns → **TRUE / FALSE** with date-format normalisation."
)
financial_cols = st.multiselect(
    "Financial columns",
    options=remaining_cols,
    help="e.g. Premium, Commission, Tax, Loss Amount.",
)

tolerance = st.number_input(
    "Tolerance (financial)",
    min_value=0.0, value=0.01, step=0.001, format="%.4f",
    help="Differences ≤ tolerance are treated as 0.",
)
st.markdown("</div>", unsafe_allow_html=True)

non_fin_cols = [c for c in remaining_cols if c not in financial_cols]
if non_fin_cols:
    st.markdown(
        "**Non-Financial columns** (TRUE/FALSE with date normalisation): "
        + ", ".join(f"`{c}`" for c in non_fin_cols)
    )

# ══════════════════════════════════════════════════════════════════════════════
# Step 5 – Generate
# ══════════════════════════════════════════════════════════════════════════════
st.markdown("---")

run_disabled = not key_cols
if not key_cols:
    st.warning("⚠️  Select at least one Key column to proceed.")

generate = st.button(
    "🚀  Generate Reconciliation",
    type="primary",
    use_container_width=True,
    disabled=run_disabled,
)

if generate:
    prog = st.progress(0, text="Initialising…")
    t_start = time.perf_counter()
    prog.progress(10, text="Running positional matching & vectorized comparison…")

    try:
        output_bytes = reconcile(
            source_df=source_df,
            target_df=target_df,
            key_cols=key_cols,
            financial_cols=financial_cols,
            tolerance=tolerance,
            col_mapping=st.session_state.get("col_mapping", {}),
        )
    except Exception as e:
        prog.empty()
        st.error(f"❌ Reconciliation failed: {e}")
        st.exception(e)
        st.stop()

    elapsed = round(time.perf_counter() - t_start, 2)
    prog.progress(90, text="Finalising workbook…")

    # Read Summary back for metric pills
    summary_vals: dict = {}
    sheet_list: list = []
    try:
        import openpyxl
        output_bytes.seek(0)
        wb_p = openpyxl.load_workbook(output_bytes, read_only=True, data_only=True)
        sheet_list = wb_p.sheetnames
        if "Summary" in sheet_list:
            for row in wb_p["Summary"].iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    summary_vals[str(row[0])] = row[1]
        wb_p.close()
        output_bytes.seek(0)
    except Exception:
        pass

    prog.progress(100, text="Done!")
    prog.empty()

    st.success(f"✅  Reconciliation complete in **{elapsed}s**!", icon="✅")
    st.markdown(
        f'<div class="timer-badge">⏱ Processing time: {elapsed} seconds</div>',
        unsafe_allow_html=True,
    )

    # Metric pills
    if summary_vals:
        METRIC_DEFS = [
            ("Total Source Records",      "Source Rows",        ""),
            ("Total Target Records",      "Target Rows",        ""),
            ("Matched Records",           "Matched",            "ok"),
            ("Records Missing in Source", "Missing in Source",  "err"),
            ("Records Missing in Target",    "Missing in Target",     "err"),
            ("Financial Differences",     "Fin. Diffs",         "warn"),
            ("Non-Financial Differences", "Non-Fin. Diffs",     "warn"),
            ("Match Percentage",          "Match %",            "ok"),
        ]
        pills_html = '<div class="metric-row">'
        for key, label, cls in METRIC_DEFS:
            val = summary_vals.get(key, "–")
            pills_html += (
                f'<div class="metric-pill {cls}">'
                f'<div class="val">{val}</div>'
                f'<div class="lbl">{label}</div>'
                "</div>"
            )
        pills_html += "</div>"
        st.markdown(pills_html, unsafe_allow_html=True)

    # Show sheet list
    if sheet_list:
        sheet_info = " · ".join(f"**{s}**" for s in sheet_list)
        st.caption(f"Workbook contains: {sheet_info}")
        if "Column_Diff" in sheet_list:
            _sf = src_only_final if src_only or tgt_only else []
            _tf = tgt_only_final if src_only or tgt_only else []
            st.info(
                f"📋 **Column_Diff** sheet included — "
                f"{len(_sf)} column(s) only in Source, "
                f"{len(_tf)} column(s) only in Target.",
                icon="⚠️",
            )

    st.markdown("---")

    st.download_button(
        label="⬇️  Download Reconciliation Workbook",
        data=output_bytes,
        file_name="Data_Reconciliation_Output.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        type="primary",
    )
